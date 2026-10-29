import streamlit as st
import pandas as pd
import docx
from docxtpl import DocxTemplate
from docxcompose.composer import Composer
from datetime import date, timedelta, datetime
import re
import io
import os
import openpyxl
from openpyxl.cell.cell import MergedCell
from copy import copy, deepcopy
import holidays
from PIL import Image as PILImage

# Nomes exatos dos arquivos que devem estar na mesma pasta
MODELO_PADRAO_NOME = "MODELOLISTA.xlsx"
MODELO_DESPERTARE_NOME = "ModeloDespertare.docx"
ARQUIVO_TUSS_PADRAO = "CODIGOS TUSS PADRÃO.txt"
ARQUIVO_LOGO = "Logo.png"

# Nova Expressão Regular flexível (Lê "17:00", "17h", "17H", "17h50" e separadores como "-", "às", "as")
PADRAO_BUSCA_HORARIO = r'(?:[0-2]?\d)(?::[0-5]\d|[hH](?:[0-5]\d)?)\s*(?:às|as|-|–|—)\s*(?:[0-2]?\d)(?::[0-5]\d|[hH](?:[0-5]\d)?)'

# ==========================================
# FUNÇÕES AUXILIARES E DE FORMATAÇÃO
# ==========================================

def carregar_tuss_padrao():
    tuss_padrao = {}
    if os.path.exists(ARQUIVO_TUSS_PADRAO):
        with open(ARQUIVO_TUSS_PADRAO, 'r', encoding='utf-8') as f:
            for linha in f:
                if '=' in linha:
                    chave, valor = linha.split('=', 1)
                    tuss_padrao[chave.strip().upper()] = valor.strip()
        
        if "ABA" in tuss_padrao and "PSICOLOGIA ABA" not in tuss_padrao:
            tuss_padrao["PSICOLOGIA ABA"] = tuss_padrao["ABA"]
            
    return tuss_padrao

def limpar_especialidade(texto):
    texto_limpo = re.sub(r'\(.*?\)', '', texto).strip()
    abreviacoes = {
        "TO": "Terapia Ocupacional",
        "FISIO": "Fisioterapia",
        "FONO": "Fonoaudiologia",
        "NUTRI": "Nutrição",
        "MUSICA": "Musicoterapia",
        "MÚSICA": "Musicoterapia",
        "ABA": "Psicologia ABA"
    }
    return abreviacoes.get(texto_limpo.upper(), texto_limpo)

def limpar_nome_arquivo(nome):
    nome_limpo = re.sub(r'[\\/*?:_\|\[\]]', '', nome)
    return re.sub(r'\s+', ' ', nome_limpo).strip()

def remover_conselho(nome_profissional):
    if not nome_profissional:
        return ""
    padrao_conselho = r'\b(CRP|CREFITO|CRM|COREN|COFFITO|OAB|CREF|CRFa|CRFª|CRN|CRT|CRBM|CBO)\b.*'
    nome_limpo = re.sub(padrao_conselho, '', nome_profissional, flags=re.IGNORECASE)
    return re.sub(r'[-\s,/|:_]+$', '', nome_limpo).strip()

def copiar_estilos_celula(origem, destino):
    if origem.font: destino.font = copy(origem.font)
    if origem.border: destino.border = copy(origem.border)
    if origem.fill: destino.fill = copy(origem.fill)
    if origem.number_format: destino.number_format = copy(origem.number_format)
    if origem.alignment: destino.alignment = copy(origem.alignment)

def limpar_style_celula(celula):
    celula.font = None
    celula.border = openpyxl.styles.Border()
    celula.fill = openpyxl.styles.PatternFill(fill_type=None)
    celula.alignment = None

def clonar_aba_com_estilo(wb_modelo, wb_destino, nome_nova_aba, imagens_cache=None, nova_posicao_imagem=None):
    aba_origem = wb_modelo.worksheets[0]
    aba_destino = wb_destino.create_sheet(title=nome_nova_aba)
    aba_destino.views.sheetView[0].showGridLines = True

    for col, col_dim in aba_origem.column_dimensions.items():
        aba_destino.column_dimensions[col].width = col_dim.width

    for row, row_dim in aba_origem.row_dimensions.items():
        aba_destino.row_dimensions[row].height = row_dim.height

    for mesclada in aba_origem.merged_cells.ranges:
        aba_destino.merge_cells(str(mesclada))

    for r in range(1, aba_origem.max_row + 1):
        for c in range(1, aba_origem.max_column + 1):
            cel_origem = aba_origem.cell(row=r, column=c)
            cel_destino = aba_destino.cell(row=r, column=c, value=cel_origem.value)
            copiar_estilos_celula(cel_origem, cel_destino)

    if imagens_cache:
        for img_data in imagens_cache:
            try:
                nova_img = openpyxl.drawing.image.Image(io.BytesIO(img_data))
                nova_img.width = int(nova_img.width * 0.7)
                nova_img.height = int(nova_img.height * 0.7)
                nova_img.anchor = nova_posicao_imagem if nova_posicao_imagem else "E25"
                aba_destino.add_image(nova_img)
            except Exception as e:
                pass

    return aba_destino

def escrever_celula_segura(aba, row, column, value):
    celula = aba.cell(row=row, column=column)
    if not isinstance(celula, MergedCell):
        celula.value = value

def processar_placeholders_e_limpar_x(aba, paciente, especialidade, profissional, cartao, senha, linha_inicio_dados):
    for r in range(1, aba.max_row + 1):
        for c in range(1, aba.max_column + 1):
            celula = aba.cell(row=r, column=c)
            if isinstance(celula, MergedCell): continue

            valor = celula.value
            if isinstance(valor, str):
                if "[PACIENTE]" in valor.upper():
                    celula.value = re.compile(re.escape("[PACIENTE]"), re.IGNORECASE).sub(paciente, valor)
                elif "[ESPECIALIDADE]" in valor.upper():
                    celula.value = re.compile(re.escape("[ESPECIALIDADE]"), re.IGNORECASE).sub(especialidade, valor)
                elif "[TERAPIA]" in valor.upper():
                    celula.value = re.compile(re.escape("[TERAPIA]"), re.IGNORECASE).sub(especialidade, valor)
                elif "[PROFISSIONAL]" in valor.upper():
                    celula.value = re.compile(re.escape("[PROFISSIONAL]"), re.IGNORECASE).sub(profissional, valor)
                elif "[TERAPEUTA]" in valor.upper():
                    celula.value = re.compile(re.escape("[TERAPEUTA]"), re.IGNORECASE).sub(profissional, valor)
                elif any(ph in valor.upper() for ph in ["[CARTÃO]", "[CARTAO]", "[CARTEIRINHA]", "[CODIGO]"]):
                    for ph in ["[CARTÃO]", "[CARTAO]", "[CARTEIRINHA]", "[CODIGO]"]:
                        if ph in valor.upper():
                            celula.value = re.compile(re.escape(ph), re.IGNORECASE).sub(cartao, valor)
                elif any(ph in valor.upper() for ph in ["[SENHA]"]):
                    for ph in ["[SENHA]"]:
                        if ph in valor.upper():
                            celula.value = re.compile(re.escape(ph), re.IGNORECASE).sub(senha, valor)
                
                elif r >= linha_inicio_dados and (valor.strip().lower() == 'x' or valor.strip() == '[]'):
                    celula.value = ""

def unificar_horarios_seguidos(horarios_lista):
    if not horarios_lista: return []
    
    def parse_tempo_para_minutos(h_str):
        # Captura H e M de forma inteligente (ex: "17:00", "17h", "17h50", "17H")
        padrao = r'([0-2]?\d)(?::([0-5]\d)|[hH]([0-5]\d)?)'
        matches = list(re.finditer(padrao, h_str))
        if len(matches) >= 2:
            try:
                h1 = int(matches[0].group(1))
                m1 = int(matches[0].group(2) or matches[0].group(3) or 0)
                h2 = int(matches[1].group(1))
                m2 = int(matches[1].group(2) or matches[1].group(3) or 0)
                # Retorna sempre uma versão limpa (HH:MM) para ficar bonito no documento final
                return (h1 * 60 + m1, h2 * 60 + m2, f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}")
            except ValueError:
                pass
        return (0, 0, "00:00", "00:00")

    intervalos = []
    for h in horarios_lista:
        min_inicio, min_fim, str_inicio, str_fim = parse_tempo_para_minutos(h)
        if str_inicio != "00:00" or str_fim != "00:00":
            intervalos.append({
                'min_ini': min_inicio, 'min_fim': min_fim,
                'str_ini': str_inicio, 'str_fim': str_fim
            })

    intervalos.sort(key=lambda x: x['min_ini'])
    unificados = []
    for item in intervalos:
        if not unificados:
            unificados.append(item)
        else:
            ultimo = unificados[-1]
            diferenca = item['min_ini'] - ultimo['min_fim']
            if 0 <= diferenca <= 15:
                ultimo['min_fim'] = item['min_fim']
                ultimo['str_fim'] = item['str_fim']
            else:
                unificados.append(item)

    return [f"{x['str_ini']} às {x['str_fim']}" for x in unificados]

def data_por_extenso(data_obj):
    if not data_obj:
        return ""
    meses = {
        1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
        5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
        9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"
    }
    return f"{data_obj.day} de {meses[data_obj.month]} de {data_obj.year}"


# ==========================================
# INTERFACE GRÁFICA STREAMLIT
# ==========================================

st.set_page_config(page_title="Gerador de Listas de Presença", layout="centered")

col_titulo, col_logo = st.columns([9, 1])
with col_titulo:
    st.title("Gerador de Listas de Presença")
with col_logo:
    if os.path.exists(ARQUIVO_LOGO):
        st.write("")
        st.write("")
        st.image(ARQUIVO_LOGO, width=100) 

tem_tuss_local = os.path.exists(ARQUIVO_TUSS_PADRAO)

tipo_plano = st.selectbox(
    "Selecione o Plano de Saúde:", 
    options=["Unimed Brasil", "Leve Saúde (Em breve)", "Assim Saúde (Em breve)", "Despertare"],
    index=0
)

# Lógica de Adaptação Dinâmica de Interface
is_docx_plan = "Despertare" in tipo_plano

if is_docx_plan:
    tem_modelo_local = os.path.exists(MODELO_DESPERTARE_NOME)
    tipo_arquivo_permitido = ["docx"]
    nome_modelo_alvo = MODELO_DESPERTARE_NOME
else:
    tem_modelo_local = os.path.exists(MODELO_PADRAO_NOME)
    tipo_arquivo_permitido = ["xlsx"]
    nome_modelo_alvo = MODELO_PADRAO_NOME

col1, col2 = st.columns(2)
mes = col1.number_input("Mês da Lista", min_value=1, max_value=12, value=datetime.now().month)
ano = col2.number_input("Ano da Lista", min_value=2020, max_value=2100, value=datetime.now().year)

senha_digitada = st.text_input("Senha da Guia (opcional):", placeholder="Digite a senha para preencher o documento" if not is_docx_plan else "Não se aplica ao modelo DOCX")

st.divider()
st.subheader("Anexe os arquivos necessários")
arquivo_grade = st.file_uploader("1. Grade de Faturamento (Word - .docx)", type=["docx"])

# ==========================================
# PRÉ-LEITURA PARA CÓDIGOS TUSS E EXTRAÇÃO
# ==========================================
especialidades_pre_encontradas = []
codigos_tuss = {}

if arquivo_grade:
    try:
        doc_temp = docx.Document(arquivo_grade)
        texto_temp = []
        for table in doc_temp.tables:
            for row in table.rows:
                linha_tabela = " | ".join([cell.text.strip().replace('\n', ' ') for cell in row.cells])
                texto_temp.append(linha_tabela)

        grade_horarios_temp = []
        for linha in texto_temp:
            if re.search(PADRAO_BUSCA_HORARIO, linha, flags=re.IGNORECASE):
                partes = [p.strip() for p in linha.split('|')]
                if len(partes) >= 6:
                    for i in range(1, 6):
                        espec = partes[i]
                        if espec:
                            espec_limpa = limpar_especialidade(espec)
                            if espec_limpa:
                                grade_horarios_temp.append(espec_limpa)

        especialidades_pre_encontradas = sorted(list(set(grade_horarios_temp)))
        arquivo_grade.seek(0) 
    except Exception:
        pass

if especialidades_pre_encontradas and not is_docx_plan:
    inserir_tuss = st.checkbox("Deseja inserir ou revisar os códigos TUSS?")
    if inserir_tuss:
        mensagem_tuss = "Digite o código TUSS correspondente."
        if tem_tuss_local:
            mensagem_tuss += " Deixe em branco para usar o padrão automático."
            
        st.write(mensagem_tuss)
        for espec in especialidades_pre_encontradas:
            codigos_tuss[espec] = st.text_input(f"[{espec}]", placeholder="Ex: 20104081")

st.divider()

arquivo_modelo = None
if tem_modelo_local:
    st.success(f"✅ O modelo padrão da clínica já está carregado para o plano {tipo_plano}.")
    arquivo_modelo = st.file_uploader(f"2. Modelo Alternativo (Opcional - anexe se quiser sobrescrever o {tipo_arquivo_permitido[0]})", type=tipo_arquivo_permitido)
else:
    st.warning(f"⚠️ O arquivo '{nome_modelo_alvo}' não foi encontrado. Por favor, anexe o modelo manualmente.")
    arquivo_modelo = st.file_uploader(f"2. Modelo da Planilha Padrão ({tipo_arquivo_permitido[0]}) [OBRIGATÓRIO]", type=tipo_arquivo_permitido)


# ==========================================
# MOTOR DE GERAÇÃO (Lógica Unificada)
# ==========================================

if st.button("Gerar Lista de Presença", type="primary"):
    if not arquivo_grade:
        st.error("Por favor, faça o upload da Grade (Word) antes de continuar.")
    elif not tem_modelo_local and not arquivo_modelo:
        st.error(f"Por favor, faça o upload do Modelo ({tipo_arquivo_permitido[0]}) antes de continuar.")
    else:
        with st.spinner('Lendo dados e gerando listas...'):
            try:
                feriados_br = holidays.Brazil(years=ano)
                
                # Leitura da Grade (Idêntica para ambos os métodos)
                doc = docx.Document(arquivo_grade)
                texto_completo = []
                for p in doc.paragraphs:
                    if p.text.strip(): texto_completo.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        linha_tabela = " | ".join([cell.text.strip().replace('\n', ' ') for cell in row.cells])
                        texto_completo.append(linha_tabela)

                paciente_cru = texto_completo[0]
                paciente_cru = re.sub(r'^paciente\s*:\s*', '', paciente_cru, flags=re.IGNORECASE)
                paciente_cru = re.sub(r'^paciente\s+', '', paciente_cru, flags=re.IGNORECASE)
                paciente = limpar_nome_arquivo(paciente_cru)

                profissionais = {}
                capturando_prof = False
                for linha in texto_completo:
                    if "PROFISSIONAIS" in linha.upper():
                        capturando_prof = True
                        continue
                    if capturando_prof and linha:
                        partes = linha.split(' ', 1)
                        if len(partes) == 2:
                            espec_chave = limpar_especialidade(partes[0].strip())
                            profissionais[espec_chave.upper()] = partes[1].strip()

                grade_horarios = []
                for linha in texto_completo:
                    if re.search(PADRAO_BUSCA_HORARIO, linha, flags=re.IGNORECASE):
                        partes = [p.strip() for p in linha.split('|')]
                        if len(partes) >= 6:
                            horario = partes[0]
                            for i in range(1, 6):
                                espec = partes[i]
                                if espec:
                                    espec_limpa = limpar_especialidade(espec)
                                    if espec_limpa:
                                        grade_horarios.append({'dia_semana': i - 1, 'horario': horario, 'especialidade': espec_limpa})

                especialidades_encontradas = sorted(list(set([item['especialidade'] for item in grade_horarios])))
                
                if not especialidades_encontradas:
                    st.error("Nenhuma especialidade compatível encontrada na grade.")
                    st.stop()


                # ==========================================
                # ROTA 1: GERAÇÃO DOCX (DESPERTARE)
                # ==========================================
                if is_docx_plan:
                    bytes_brutos_modelo = arquivo_modelo.getvalue() if arquivo_modelo else open(MODELO_DESPERTARE_NOME, "rb").read()
                    
                    sessoes_por_especialidade = {}
                    todas_as_datas_encontradas = []

                    for espec in especialidades_encontradas:
                        dias_dessa_espec = {}
                        for item in grade_horarios:
                            if item['especialidade'] == espec:
                                dia_sem = item['dia_semana']
                                if dia_sem not in dias_dessa_espec:
                                    dias_dessa_espec[dia_sem] = []
                                dias_dessa_espec[dia_sem].append(item['horario'])

                        for dia_sem in dias_dessa_espec:
                            dias_dessa_espec[dia_sem] = unificar_horarios_seguidos(dias_dessa_espec[dia_sem])

                        lista_sessoes = []
                        data_atual = date(ano, mes, 1)
                        while data_atual.month == mes:
                            if data_atual.weekday() in dias_dessa_espec:
                                if data_atual.weekday() in [5, 6] or data_atual in feriados_br:
                                    data_atual += timedelta(days=1)
                                    continue
                                for horario_sessao in dias_dessa_espec[data_atual.weekday()]:
                                    lista_sessoes.append({
                                        "data_obj": data_atual,
                                        "data": data_atual.strftime("%d/%m/%Y"),
                                        "horario": horario_sessao
                                    })
                                    todas_as_datas_encontradas.append(data_atual)
                            data_atual += timedelta(days=1)

                        sessoes_por_especialidade[espec] = lista_sessoes

                    primeira_data_global_obj = min(todas_as_datas_encontradas) if todas_as_datas_encontradas else None
                    primeira_data_global_extenso = data_por_extenso(primeira_data_global_obj)

                    buffers_gerados = []

                    for espec in especialidades_encontradas:
                        nome_profissional_completo = profissionais.get(espec.upper(), "Profissional não cadastrado")
                        lista_sessoes = sessoes_por_especialidade[espec]

                        doc_temp = docx.Document(io.BytesIO(bytes_brutos_modelo))
                        tabela_alvo = None

                        for tb in doc_temp.tables:
                            if len(tb.rows) > 0:
                                texto_cabecalho = " ".join([c.text.upper() for c in tb.rows[0].cells])
                                if "DATA" in texto_cabecalho and "HORÁRIO" in texto_cabecalho:
                                    tabela_alvo = tb
                                    break

                        linhas_adicionadas = max(0, len(lista_sessoes) - 1)

                        if tabela_alvo and lista_sessoes:
                            linha_molde = tabela_alvo.rows[1]

                            for cell in linha_molde.cells:
                                cell.text = ""

                            for _ in range(linhas_adicionadas):
                                nova_linha = deepcopy(linha_molde._tr)
                                tabela_alvo._tbl.append(nova_linha)

                            for idx, sessao in enumerate(lista_sessoes):
                                linha_atual = tabela_alvo.rows[idx + 1]
                                linha_atual.cells[0].text = sessao['data']
                                linha_atual.cells[1].text = sessao['horario']
                                linha_atual.cells[2].text = ""

                        buffer_tabela = io.BytesIO()
                        doc_temp.save(buffer_tabela)
                        buffer_tabela.seek(0)

                        doc_final = DocxTemplate(buffer_tabela)
                        contexto = {
                            "paciente": paciente,
                            "cpf": "Não consta na grade",
                            "nascimento": "Não consta na grade",
                            "primeira_data": primeira_data_global_extenso,
                            "profissional": nome_profissional_completo
                        }
                        doc_final.render(contexto)

                        buffer_pos_render = io.BytesIO()
                        doc_final.save(buffer_pos_render)
                        buffer_pos_render.seek(0)

                        doc_limpo = docx.Document(buffer_pos_render)

                        # REMOÇÃO CIRÚRGICA DE ESPAÇOS INTERMEDIÁRIOS
                        idx_data_extenso = -1
                        idx_assinatura = -1

                        for idx, p in enumerate(doc_limpo.paragraphs):
                            if primeira_data_global_extenso and primeira_data_global_extenso in p.text:
                                idx_data_extenso = idx
                            if "____" in p.text or (nome_profissional_completo in p.text and idx_assinatura == -1):
                                idx_assinatura = idx

                        if idx_data_extenso != -1 and idx_assinatura != -1 and idx_assinatura > idx_data_extenso:
                            paragrafos_vazios_intervalo = []
                            for idx in range(idx_data_extenso + 1, idx_assinatura):
                                p = doc_limpo.paragraphs[idx]
                                if not p.text.strip():
                                    paragrafos_vazios_intervalo.append(p)

                            LIMITE_SESSOES_SEGURAS = 6
                            if len(lista_sessoes) > LIMITE_SESSOES_SEGURAS:
                                excesso = len(lista_sessoes) - LIMITE_SESSOES_SEGURAS
                                remover_qtd = min(len(paragrafos_vazios_intervalo), excesso)
                                for i in range(remover_qtd):
                                    p_remover = paragrafos_vazios_intervalo[i]
                                    try:
                                        p_remover._element.getparent().remove(p_remover._element)
                                    except Exception:
                                        pass

                        buffer_memoria = io.BytesIO()
                        doc_limpo.save(buffer_memoria)
                        buffer_memoria.seek(0)
                        buffers_gerados.append(buffer_memoria)

                    # Costura final em memória
                    if buffers_gerados:
                        master_doc = docx.Document(buffers_gerados[0])
                        composer = Composer(master_doc)

                        for buffer_extra in buffers_gerados[1:]:
                            master_doc.add_page_break()
                            doc_adicional = docx.Document(buffer_extra)
                            composer.append(doc_adicional)

                        buffer_final = io.BytesIO()
                        composer.save(buffer_final)
                        buffer_final.seek(0)

                        nome_arquivo_final = f"{paciente} Listas Despertare.docx"
                        st.success("Listas geradas com sucesso!")
                        st.download_button(
                            label="⬇️ Baixar Lista Pronta (Word)",
                            data=buffer_final,
                            file_name=nome_arquivo_final,
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            type="primary"
                        )

                # ==========================================
                # ROTA 2: GERAÇÃO EXCEL (UNIMED, ETC)
                # ==========================================
                else:
                    tuss_padrao_dit = carregar_tuss_padrao()
                    
                    if arquivo_modelo is not None:
                        wb_modelo = openpyxl.load_workbook(arquivo_modelo, data_only=False)
                    else:
                        wb_modelo = openpyxl.load_workbook(MODELO_PADRAO_NOME, data_only=False)
                    
                    aba_modelo_original = wb_modelo.worksheets[0]

                    imagens_cache = []
                    if hasattr(aba_modelo_original, '_images') and aba_modelo_original._images:
                        for img in aba_modelo_original._images:
                            try:
                                imagens_cache.append(img._data())
                            except: pass

                    linha_inicio_dados = 9
                    for r in range(1, aba_modelo_original.max_row + 1):
                        cel_valor = aba_modelo_original.cell(row=r, column=1).value
                        if isinstance(cel_valor, str) and (cel_valor.strip() == '[]' or cel_valor.strip().lower() == 'x'):
                            linha_inicio_dados = r
                            break

                    cartao = ""
                    padrao_cartao = re.compile(r'\b(cart(?:ã|a)o|carteirinha|cod(?:igo)?\s*do?\s*cart(?:ã|a)o|n(?:u|o|º)?\s*da?\s*carteirinha)\s*[:-]?\s*(.*)', re.IGNORECASE)
                    for linha in texto_completo:
                        match_cartao = padrao_cartao.search(linha)
                        if match_cartao:
                            cartao = match_cartao.group(2).strip()
                            break

                    wb_novo = openpyxl.Workbook()
                    default_sheet = wb_novo.active

                    for espec in especialidades_encontradas:
                        nome_profissional_completo = profissionais.get(espec.upper(), "Profissional não cadastrado")
                        nome_profissional_cabecalho = remover_conselho(nome_profissional_completo)

                        dias_dessa_espec = {}
                        for item in grade_horarios:
                            if item['especialidade'] == espec:
                                dia_sem = item['dia_semana']
                                if dia_sem not in dias_dessa_espec:
                                    dias_dessa_espec[dia_sem] = []
                                dias_dessa_espec[dia_sem].append(item['horario'])

                        for dia_sem in dias_dessa_espec:
                            dias_dessa_espec[dia_sem] = unificar_horarios_seguidos(dias_dessa_espec[dia_sem])

                        codigo_digitado = codigos_tuss.get(espec, "").strip()
                        codigo_final = codigo_digitado if codigo_digitado else tuss_padrao_dit.get(espec.upper(), "")

                        datas_mes = []
                        data_atual = date(ano, mes, 1)
                        while data_atual.month == mes:
                            if data_atual.weekday() in dias_dessa_espec:
                                if data_atual.weekday() in [5, 6] or data_atual in feriados_br:
                                    data_atual += timedelta(days=1)
                                    continue

                                for horario_sessao in dias_dessa_espec[data_atual.weekday()]:
                                    datas_mes.append({
                                        "Data": data_atual.strftime("%d/%m/%Y"),
                                        "Horário": horario_sessao,
                                        "Código TUSS": codigo_final,
                                        "Profissional": nome_profissional_completo,
                                        "Assinatura": ""
                                    })
                            data_atual += timedelta(days=1)

                        nome_aba = limpar_nome_arquivo(espec)[:30] or "Terapia"
                        limite_preenchido = linha_inicio_dados + len(datas_mes)
                        
                        aba_pronta = clonar_aba_com_estilo(
                            wb_modelo, wb_novo, nome_aba, 
                            imagens_cache=imagens_cache, 
                            nova_posicao_imagem=f"E{limite_preenchido + 1}"
                        )

                        processar_placeholders_e_limpar_x(
                            aba_pronta, paciente, espec, nome_profissional_cabecalho, 
                            cartao, senha_digitada, linha_inicio_dados
                        )

                        estilos_colunas = {col: aba_modelo_original.cell(row=linha_inicio_dados, column=col) for col in range(1, 6)}

                        for idx, registro in enumerate(datas_mes):
                            linha_atual = linha_inicio_dados + idx
                            escrever_celula_segura(aba_pronta, linha_atual, 1, registro["Data"])
                            escrever_celula_segura(aba_pronta, linha_atual, 2, registro["Horário"])
                            escrever_celula_segura(aba_pronta, linha_atual, 3, registro["Código TUSS"])
                            escrever_celula_segura(aba_pronta, linha_atual, 4, registro["Profissional"])
                            escrever_celula_segura(aba_pronta, linha_atual, 5, registro["Assinatura"])

                            for col in range(1, 6):
                                celula_atual = aba_pronta.cell(row=linha_atual, column=col)
                                if not isinstance(celula_atual, MergedCell):
                                    copiar_estilos_celula(estilos_colunas[col], celula_atual)

                        for r in range(limite_preenchido, aba_pronta.max_row + 1):
                            for col in range(1, 6):
                                celula_extrapolada = aba_pronta.cell(row=r, column=col)
                                if not isinstance(celula_extrapolada, MergedCell):
                                    if celula_extrapolada.value in [None, "", "x", "[]"]:
                                        celula_extrapolada.value = ""
                                        limpar_style_celula(celula_extrapolada)

                    if default_sheet in wb_novo.worksheets and len(wb_novo.worksheets) > 1:
                        wb_novo.remove(default_sheet)
                    wb_novo.active = 0

                    buffer = io.BytesIO()
                    wb_novo.save(buffer)
                    buffer.seek(0)
                    
                    nome_arquivo_final = f"{paciente} Listas de Presenca.xlsx"

                    st.success("Listas geradas com sucesso!")
                    st.download_button(
                        label="⬇️ Baixar Lista Pronta (Excel)",
                        data=buffer,
                        file_name=nome_arquivo_final,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )

            except Exception as e:
                st.error(f"Ocorreu um erro ao processar os arquivos: {e}")
